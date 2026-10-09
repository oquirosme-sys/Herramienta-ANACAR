"""Construye la plantilla 'ANACAR Web.xlsx' (hojas, fórmulas, catálogos y formatos).
Las macros VBA se agregan después con importar.py (Excel por COM) y se guarda como .xlsm."""
import json, sys, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.drawing.image import Image as XLImage

SEED = sys.argv[1]; OUT = sys.argv[2]; LOGO = sys.argv[3]
s = open(SEED, encoding="utf8").read()
D = json.loads(s[s.index("{"):s.rindex("}") + 1])

wb = Workbook()
FN = "Century Gothic"
VERDE = PatternFill("solid", fgColor="E6F4E1")
GRIS = PatternFill("solid", fgColor="D9D9D9")
GRIS2 = PatternFill("solid", fgColor="EDEDED")
AZUL = PatternFill("solid", fgColor="DDEBF7")
fino = Side(style="thin", color="000000"); medio = Side(style="medium", color="000000")
BORDE = Border(left=fino, right=fino, top=fino, bottom=fino)
BORDEM = Border(left=medio, right=medio, top=medio, bottom=medio)
CENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
IZQ = Alignment(horizontal="left", vertical="center", wrap_text=False)


def nombre(n, ref):
    wb.defined_names[n] = DefinedName(n, attr_text=ref)


def f(ws, cel, val, bold=False, size=10, fill=None, color=None, al=None, borde=False, fmt=None):
    c = ws[cel]; c.value = val
    c.font = Font(name=FN, bold=bold, size=size, color=color)
    if fill: c.fill = fill
    if al: c.alignment = al
    if borde: c.border = BORDE
    if fmt: c.number_format = fmt
    return c


def cal(v):
    if v is None: return ""
    if isinstance(v, float) and v.is_integer(): v = int(v)
    return str(v)


# ======================= catálogos =======================
def hoja_tabla(nombre_hoja, cab, filas):
    ws = wb.create_sheet(nombre_hoja)
    for j, t in enumerate(cab): f(ws, L(j + 1) + "1", t, bold=True, fill=GRIS)
    for i, fila in enumerate(filas):
        for j, v in enumerate(fila): ws.cell(2 + i, 1 + j, v)
    for j in range(len(cab)): ws.column_dimensions[L(j + 1)].width = 16
    ws.freeze_panes = "A2"
    return ws


ws = hoja_tabla("CAT_DET", ["Detalle", "Tipo", "Descripción", "V", "Fases", "Hilos", "F. demanda", "F. diversidad", "FP", "Continua", "Mult.", "Fundamento"],
                [[d["id"], d["tipo"], d["descripcion"], d["v"], d["fases"], d["hilos"], d["fd"], d["fdiv"], d["fp"], "SI" if d.get("continua") else "NO", d.get("mult", 1.25), d.get("fundamento", "")] for d in D["detallesCarga"]])
nombre("T_DET", "CAT_DET!$A$2:$L$300"); nombre("L_DET", "CAT_DET!$A$2:$A$300")
ws = hoja_tabla("CAT_TIPOS", ["Código", "Nombre", "Método", "F. demanda", "F. diversidad"], [[t["id"], t["nombre"], t["metodo"], t["fd"], t["fdiv"]] for t in D["tiposCarga"]])
nombre("T_TIPOS", "CAT_TIPOS!$A$2:$E$40")

# conductores y tablas NEC (cada tabla con su nombre)
wc = wb.create_sheet("CAT_COND")
def bloque(col, titulo, cab, filas, nom, nomcols=None):
    c0 = CI(col)
    f(wc, L(c0) + "1", titulo, bold=True, fill=GRIS)
    for j, t in enumerate(cab): f(wc, L(c0 + j) + "2", t, bold=True, fill=GRIS2)
    for i, fila in enumerate(filas):
        for j, v in enumerate(fila): wc.cell(3 + i, c0 + j, v)
    ult = 2 + len(filas)
    nombre(nom, "CAT_COND!$%s$3:$%s$%d" % (L(c0), L(c0 + len(cab) - 1), ult))
    return c0, ult
bloque("A", "Calibre por ampacidad CU", ["Desde (A)", "Calibre"], [[x["amp"], cal(x["cal"])] for x in D["ampCU"]], "T_AMPCU")
bloque("D", "Calibre por ampacidad AL", ["Desde (A)", "Calibre"], [[x["amp"], cal(x["cal"])] for x in D["ampAL"]], "T_AMPAL")
bloque("G", "Protección por ampacidad", ["Desde (A)", "Protección (A)"], [[x["amp"], x["prot"]] for x in D["protecciones"]], "T_PROT")
c0, ult = bloque("J", "Tierra 250.122", ["Hasta (A)", "CU", "AL", "CU-MC", "AL-MC"], [[x["amp"], cal(x["CU"]), cal(x["AL"]), cal(x["CU-MC"]), cal(x["AL-MC"])] for x in D["tierras"]], "T_TIERRA_ALL")
nombre("T_TIERRA", "CAT_COND!$K$3:$N$%d" % ult); nombre("T_TIERRA_A", "CAT_COND!$J$3:$J$%d" % ult); nombre("T_TIERRA_H", "CAT_COND!$K$2:$N$2")
bloque("P", "Tubería C.10 (mm)", ["Calibre", "THHN", "RHW", "XHHW"], [[cal(x["cal"]), x["THHN"], x["RHW"], x["XHHW"]] for x in D["conduit"]], "T_CONDUIT")
c0, ult = bloque("U", "Electrodo 250.66", ["Calibre CU", "CU", "CU-MC", "Calibre AL", "AL", "AL-MC"], [[cal(x["calCU"]), cal(x["CU"]), cal(x["CU-MC"]), cal(x["calAL"]), cal(x["AL"]), cal(x["AL-MC"])] for x in D["electrodo"]], "T_ELEC")
nombre("T_ELEC_CALCU", "CAT_COND!$U$3:$U$%d" % ult); nombre("T_ELEC_CU", "CAT_COND!$V$3:$V$%d" % ult); nombre("T_ELEC_CUMC", "CAT_COND!$W$3:$W$%d" % ult)
nombre("T_ELEC_CALAL", "CAT_COND!$X$3:$X$%d" % ult); nombre("T_ELEC_AL", "CAT_COND!$Y$3:$Y$%d" % ult); nombre("T_ELEC_ALMC", "CAT_COND!$Z$3:$Z$%d" % ult)
# FAC (4 filas: FP 1, 0.95, 0.9, 0.8) — conduit magnético, como el Excel
hdr = ["1000", "900", "800", "750", "700", "600", "500", "400", "350", "300", "250", "4/0", "3/0", "2/0", "1/0", "2", "4", "6", "8", "10", "12", "14"]
f(wc, "AB1", "FAC caída de tensión (conduit magnético)", bold=True, fill=GRIS)
for j, h in enumerate(hdr): wc.cell(2, CI("AC") + j, h)
for k, (tabla, r0) in enumerate((("CU_mag", 3), ("AL_mag", 8))):
    if k: [wc.cell(7, CI("AC") + j, h) for j, h in enumerate(hdr)]
    for i, fp in enumerate(["1", "0.95", "0.9", "0.8"]):
        wc.cell(r0 + i, CI("AB"), float(fp))
        for j, h in enumerate(hdr): wc.cell(r0 + i, CI("AC") + j, D["fac"][tabla].get(fp, {}).get(h))
nombre("T_FAC_H", "CAT_COND!$AC$2:$AX$2"); nombre("T_FAC_CU", "CAT_COND!$AC$3:$AX$6"); nombre("T_FAC_AL", "CAT_COND!$AC$8:$AX$11")
keys = ["CU|60", "CU|75", "CU|90", "CU-MC|60", "CU-MC|75", "CU-MC|90", "AL|60", "AL|75", "AL|90", "AL-MC|60", "AL-MC|75", "AL-MC|90"]
bloque("BA", "Ampacidad 310.16", ["Calibre"] + keys, [[cal(x["cal"])] + [x.get(k) for k in keys] for x in D["ampacidad31016"]], "T_AMP")
bloque("BO", "Factor de temperatura 310.15(B)", ["Ambiente", "60", "75", "90"], [[k, v["60"], v["75"], v["90"]] for k, v in D["tempFactor"].items()], "T_TEMP")
bloque("BT", "Agrupamiento 310.15(C)(1)", ["Conductores", "Factor"], [[k, v] for k, v in D["agrupamiento"].items()], "T_AGRUP")
ck = ["THHN|PCV SCH40", "THHN|EMT", "RHW|PCV SCH40", "RHW|EMT", "XHHW-2|PCV SCH40", "XHHW-2|EMT", "DUCTOBARRA|DUCTOBARRA"]
bloque("BW", "Constantes C (cortocircuito)", ["Calibre"] + ck, [[cal(x["cal"])] + [x.get(k) for k in ck] for x in D["constC"]], "T_CONSTC")
listas = {"CH": ("Sistemas", D["listas"]["sistemas"], "L_SISTEMAS"), "CI": ("Materiales", D["listas"]["materiales"], "L_MATERIALES"), "CJ": ("Aislamientos", D["aislamientos"], "L_AISLAMIENTOS"),
          "CK": ("Tuberías", D["listas"]["tuberias"], "L_TUBERIAS"), "CL": ("Montajes", D["listas"]["montajes"], "L_MONTAJES"), "CM": ("Temperaturas", D["listas"]["temperaturas"], "L_TEMPS"),
          "CN": ("Agrupamientos", D["listas"]["agrupamientos"], "L_AGRUP"), "CO": ("Marcas", sorted({b["marca"] for b in D["breakers"]}), "L_MARCAS"),
          "CP": ("Unidades", ["STD", "GFCI", "AFCI", "AFCI/GFCI", "HACR", "SHT"], "L_UNIDADES"), "CQ": ("Disparo principal", ["TM", "LI", "LSI", "LSIG"], "L_DISPARO"),
          "CR": ("SI/NO", ["SI", "NO"], "L_SINO"), "CS": ("Tipo", ["3F", "1F"], "L_TIPO"), "CT": ("Principal", ["Interruptor", "Zapatas"], "L_PRINCIPAL"),
          "CU": ("Tipo de equipo", ["Tablero", "Subestación"], "L_CLASE"), "CV": ("Transferencia", ["ATS", "MTS", "IP"], "L_TRANSF"), "CW": ("Fuente alterna", ["Generador", "Tablero"], "L_FUENTE"),
          "CX": ("Ocupación", ["comercial", "vivienda", "hotel", "hospital", "industrial"], "L_OCUP"), "CY": ("Carga derivados", ["conectada", "demandada"], "L_CDERIV")}
for col, (t, vals, nm) in listas.items():
    f(wc, col + "1", t, bold=True, fill=GRIS)
    for i, v in enumerate(vals): wc[col + str(2 + i)] = v
    nombre(nm, "CAT_COND!$%s$2:$%s$%d" % (col, col, 1 + len(vals)))
if "marcas" not in D: pass

hoja_tabla("CAT_BREAKERS", ["Marca", "Tipo", "Modelo", "Marco (A)", "Amperios", "Unidad", "Polos", "SCCR (kA)", "SCCR a (V)", "Nota"],
           [[b["marca"], b["id"], b.get("modelo") or "", b.get("marco"), b.get("amperios"), b.get("unidad"), b.get("polos"), b.get("sccr"), b.get("vSccr"), b.get("nota") or ""] for b in D["breakers"]])
hoja_tabla("CAT_TABLEROS", ["Tipo", "Modelo", "Fabricante", "Barra fase", "Barra neutro", "Barra tierra", "Espacios", "Fases", "Principal", "Nota"],
           [[t["id"], t["modelo"], t["fabricante"], t["barraFase"], t["barraNeutro"], t["barraTierra"], t["espacios"], t.get("fases", "3F"), t.get("principal", "Ambos"), t.get("nota") or ""] for t in D["tablerosCat"]])
hoja_tabla("CAT_SPD", ["Marca", "Tipo", "Modelo", "Montaje", "kA L-L", "kA L-N", "Voltaje", "Fases", "Conexión"],
           [[x["marca"], x["id"], x["modelo"], x["montaje"], x["kaLL"], x["kaLN"], x["voltaje"], x["fases"], x.get("conexion") or ""] for x in D["supresores"]])
hoja_tabla("CAT_TRAFOS", ["#", "Nombre", "Tipo", "KACC 480", "KACC 240", "KACC 208", "Z %"], [[t["id"], t["nombre"], t["tipo"], t["kacc480"], t["kacc240"], t["kacc208"], t["z"]] for t in D["transformadores"]])
nombre("L_TRAFOS", "CAT_TRAFOS!$B$2:$B$100")


def a_like(pat):
    """Expresión regular del catálogo web → patrones Like de VBA separados por coma."""
    if pat == "^$": return "(vacio)"
    p = pat.lstrip("^")
    m = re.match(r"^\((.*)\)$", p)
    if m: return ",".join(x + "*" for x in m.group(1).split("|"))
    if p.endswith("$"): return p[:-1]
    return p + "*"
hoja_tabla("CAT_REGLAS_T", ["Marca", "Familia", "Patrones del modelo (Like)", "V máx.", "Fases", "Breakers ramales", "Breakers principales", "Preferencia", "Clase", "Nota"],
           [[r["marca"], r["familia"], a_like(r["patron"]), r["vMax"], r["fases"], r["ramales"], r["principales"], r["orden"], r.get("clase", "tablero"), r.get("nota", "")] for r in D["reglas"]["tableros"]])
hoja_tabla("CAT_REGLAS_B", ["Marca", "Familia", "Patrones del modelo (Like)", "V máx.", "Plantilla"],
           [[r["marca"], r["familia"], a_like(r["patron"]), r["vMax"], r.get("plantilla", "")] for r in D["reglas"]["breakers"]])
hoja_tabla("CAT_SERIE", ["Marca", "Voltaje", "kA serie", "Principal máx. (A)", "Línea (principal)", "Carga (ramales)", "Nota"],
           [[c.get("marca", "Eaton"), c["voltaje"], c["kA"], c.get("principalMax"), c["linea"], c["carga"], c.get("nota", "")] for c in D["seriesRating"]["combinaciones"]])
hoja_tabla("CAT_USOS", ["#", "Tipo de uso", "Iluminación VA/m²", "Tomas VA/m²", "Climatización VA/m²", "Nota"], [[u["id"], u["nombre"], u["ilum"], u["tomas"], u["clima"], u["nota"]] for u in D["usos"]])
nombre("T_USOS", "CAT_USOS!$B$2:$E$100"); nombre("L_USOS", "CAT_USOS!$B$2:$B$100")

# ======================= MC_MACHOTE: memoria de cálculo de un tablero =======================
mc = wb.create_sheet("MC_MACHOTE", 0)
F1, R0, RN = 25, 25, 124
f(mc, "B1", "MEMORIA DE CÁLCULO — TABLERO", bold=True, size=16, color="006600")
mc["E1"] = '=$BX$2'; mc["E1"].font = Font(name=FN, bold=True, size=16, color="006600")
entradas = [  # (columna etiqueta, columna valor, fila, etiqueta, valor inicial, validación)
    ("B", "C", 3, "Nombre / ID", "MACHOTE", None), ("B", "C", 4, "Prefijo", "TABLERO", None), ("B", "C", 5, "Tipo", "3F", "L_TIPO"), ("B", "C", 6, "Sistema (V)", "120/208", "L_SISTEMAS"),
    ("B", "C", 10, "Alimentado desde", None, "LISTA_TABLEROS"), ("B", "C", 11, "Conectado a (si es acometida)", "", None), ("B", "C", 12, "Longitud alimentador (m)", 10, None),
    ("B", "C", 13, "Reserva (fracción)", 0.1, None), ("B", "C", 14, "Icc disponible manual (kA)", None, None), ("B", "C", 15, "Marca (vacío = proyecto)", None, "L_MARCAS"),
    ("B", "C", 16, "Principal", "Interruptor", "L_PRINCIPAL"), ("B", "C", 17, "Espacios (vacío = auto)", None, None), ("B", "C", 18, "Montaje", "Superficial", "L_MONTAJES"),
    ("B", "C", 19, "F. diversidad del tablero", 1, None), ("B", "C", 20, "Tipo de equipo", "Tablero", "L_CLASE"),
    ("E", "F", 3, "Material", "CU", "L_MATERIALES"), ("E", "F", 4, "Aislamiento", "XHHW-2", "L_AISLAMIENTOS"), ("E", "F", 5, "Tubería", "EMT", "L_TUBERIAS"),
    ("E", "F", 6, "# en paralelo (vacío = auto)", None, None), ("E", "F", 7, "Aumento de calibre (%)", 0, None), ("E", "F", 8, "Factor de potencia", 0.9, None),
    ("E", "F", 9, "Factor de protección (%)", 125, None), ("E", "F", 10, "Protección manual (A)", None, None), ("E", "F", 11, "Temp. ambiente (°C)", "26-30", "L_TEMPS"),
    ("E", "F", 12, "Temp. bornes (°C)", 75, None), ("E", "F", 13, "Agrupamiento (vacío = auto)", None, "L_AGRUP"), ("E", "F", 14, "Unidad de disparo", "TM", "L_DISPARO"),
    ("E", "F", 15, "100 % rated", "NO", "L_SINO"), ("E", "F", 16, "Breaker principal manual (tipo)", None, None), ("E", "F", 17, "Tablero catálogo manual (tipo)", None, None),
    ("E", "F", 18, "SPD manual (tipo)", None, None), ("E", "F", 19, "Calibre fijo (no aumentar)", "NO", "L_SINO"),
    ("H", "I", 3, "Transformador aguas arriba", "NO", "L_SINO"), ("H", "I", 4, "Transformador kVA", None, None), ("H", "I", 5, "Impedancia Z (%)", None, None),
    ("H", "I", 6, "Relación X/R (vacío = 4)", None, None), ("H", "I", 7, "Voltaje primario (V)", None, None), ("H", "I", 8, "Nombre transformador", "", None),
    ("H", "I", 10, "Alimentado por UPS", "NO", "L_SINO"), ("H", "I", 11, "UPS kVA", None, None), ("H", "I", 12, "Icc del inversor (× In)", 2, None),
    ("H", "I", 14, "Segunda acometida", "NO", "L_SINO"), ("H", "I", 15, "Equipo de transferencia", "ATS", "L_TRANSF"), ("H", "I", 16, "Fuente alterna", "Generador", "L_FUENTE"),
    ("H", "I", 17, "Tablero de origen (bypass)", None, "LISTA_TABLEROS"), ("H", "I", 18, "Generador kVA", None, None), ("H", "I", 19, "X''d del generador (%)", 12, None),
    ("H", "I", 20, "Longitud segunda acometida (m)", None, None), ("H", "I", 21, "SCCR del equipo (kA)", None, None)]
dvs = {}
for le, ve, fila, et, val, lista in entradas:
    f(mc, "%s%d" % (le, fila), et, size=9, al=IZQ)
    c = f(mc, "%s%d" % (ve, fila), val, size=10, fill=VERDE, borde=True, al=Alignment(horizontal="center"))
    if lista:
        if lista not in dvs:
            dvs[lista] = DataValidation(type="list", formula1="=" + lista, allow_blank=True); mc.add_data_validation(dvs[lista])
        dvs[lista].add("%s%d" % (ve, fila))
for fila, et, form in [(7, "Voltaje nominal (V)", '=IFERROR(MAX(VALUE(LEFT($C$6,FIND("/",$C$6)-1)),VALUE(MID($C$6,FIND("/",$C$6)+1,9))),N(VALUE($C$6)))'),
                       (8, "Fases", '=IF($C$5="3F",3,IF(ISNUMBER(FIND("/",$C$6)),2,1))'), (9, "Hilos", '=IF($C$5="3F",4,IF($C$8=2,3,2))')]:
    f(mc, "B%d" % fila, et, size=9); f(mc, "C%d" % fila, form, fill=GRIS2, borde=True, al=Alignment(horizontal="center"))
for c in "BEH": mc.column_dimensions[c].width = 30
mc.column_dimensions["C"].width = 16; mc.column_dimensions["F"].width = 14; mc.column_dimensions["I"].width = 14

# resultados del tablero (columna BX) — los usan las otras hojas
BX = {}
def rx(fila, et, form):
    BX[fila] = form
    mc["BW%d" % fila] = et; mc["BX%d" % fila] = form
V, FA, HI = "$C$7", "$C$8", "$C$9"
def corr_alim(x):
    return ('IF({x}=0,0,IF(AND({V}=120,{F}=1),{x}*1000/{V},IF(AND({V}=208,{F}=1),{x}*1000/{V},IF(AND({V}=240,{F}=2),{x}*1000/({V}/2),'
            'IF(AND({V}=230,{F}=1),{x}*1000/{V},IF(AND({V}=240,{F}=1),{x}*1000/({V}/2),IF(AND({V}=277,{F}=1),{x}*1000/{V},{x}*1000/({V}/1.732))))))))').format(x=x, V=V, F=FA)
MAT = 'IF($F$3="","CU",$F$3)'; AIS = 'IF($F$4="","XHHW-2",$F$4)'; TUB = 'IF($F$5="","EMT",$F$5)'
ESAL = 'OR($F$3="AL",$F$3="AL-MC")'
rx(2, "Tablero", '=TRIM($C$4&" "&$C$3)')
rx(3, "kVA conectados sin reserva [J116]", '=SUMIFS($X$25:$X$124,$T$25:$T$124,"")')
rx(4, "kVA fase A", '=SUMIFS($AJ$25:$AJ$124,$T$25:$T$124,"")'); rx(5, "kVA fase B", '=SUMIFS($AK$25:$AK$124,$T$25:$T$124,"")'); rx(6, "kVA fase C", '=SUMIFS($AL$25:$AL$124,$T$25:$T$124,"")')
rx(7, "% desbalance", '=IFERROR(ROUND(IF($C$8=3,(MAX(BX4:BX6)-MIN(BX4:BX6))/MAX(BX4:BX6)*100,IF($C$8=2,(MAX(BX4:BX5)-MIN(BX4:BX5))/MAX(BX4:BX5)*100,0)),2),0)')
rx(8, "kVA conectados [W128]", '=$F$146'); rx(9, "kVA reserva [W129]", '=$G$146'); rx(10, "kVA totales [W130]", '=$H$146'); rx(11, "kVA demandados [J134]", '=$L$146')
rx(12, "Factor de demanda [R139]", '=IF(BX10=0,1,BX11/BX10)'); rx(13, "F. diversidad [S139]", '=IF(N($C$19)>0,$C$19,1)'); rx(14, "kVA demandados alimentador [L139]", '=BX10*BX12/BX13')
rx(15, "Factor de potencia [T139]", '=IF(N($F$8)>0,$F$8,0.9)')
for i, fila in enumerate((16, 17, 18)): rx(fila, "kVA dem. fase " + "ABC"[i], '=IF($BX$3=0,0,$BX$14*BX%d/$BX$3)' % (4 + i))
for i, fila in enumerate((19, 20, 21)): rx(fila, "Corriente fase " + "ABC"[i] + " (A)", "=" + corr_alim("BX%d" % (16 + i)))
rx(22, "Factor de protección [AD139]", '=IF($F$15="SI",1,IF(N($F$9)>0,$F$9/100,1.25))')
rx(23, "Ampacidad requerida [AE139]", '=IF(BX10=0,"",IF($C$8=3,MAX(BX19:BX21),IF(BX4>0,BX19,IF(BX5>0,BX20,BX21)))*BX22)')
rx(24, "Protección (A) [AF139]", '=IF(BX23="","",IF(N($F$10)>0,$F$10,VLOOKUP(BX23,T_PROT,2,TRUE)))')
rx(25, "Paralelos base", '=IF(BX23="",1,IF(N($F$6)>0,$F$6,IF(BX23>1600,5,IF(BX23>1000,4,IF(BX23>600,3,IF(BX23>300,2,1))))))')
rx(26, "Aumento (factor)", '=1+N($F$7)/100')
rx(27, "Calibre base", '=IF(BX24="","",VLOOKUP(BX24*BX26/BX25,CHOOSE(IF(%s,2,1),T_AMPCU,T_AMPAL),2,TRUE))' % ESAL)
rx(28, "Agrupamiento", '=IF($F$13<>"",$F$13,IF(AND($C$8=3,$C$9=4),"4-6","1-3"))')
rx(29, "Validación 310.15 (n|cal|ok|cap|amp|ft|fg|term|corr)", '=IF(BX27="","",CAL_AJUSTE(BX27,%s,BX25,%s,IF($F$11="","26-30",$F$11),IF(N($F$12)>0,$F$12,75),BX28,BX23,BX24,AND(AUTO_AMP="SI",$F$19<>"SI"),BX26))' % (MAT, AIS))
rx(30, "# en paralelo", '=IF(BX29="",BX25,PARTEN(BX29,1))'); rx(31, "Calibre", '=IF(BX29="",BX27,PARTE(BX29,2))')
rx(32, "Prefijo fases", '=IF(BX30=1,"",BX30&"x")&IF($C$8=3,"3#",IF(AND($C$7>=208,$C$7<=240),"2#",""))')
rx(33, "Neutro", '=IF(BX31="","",IF($C$9=0,"",IF(AND($C$7>=208,$C$8=3,$C$9=4),BX31,IF(AND($C$8=3,$C$9=3),"",IF(AND($C$7>=208,$C$8=1,$C$9=3),"",BX31)))))')
rx(34, "Tierra 250.122", '=IF(BX24="","",INDEX(T_TIERRA,MATCH(TRUE,INDEX(T_TIERRA_A>=BX24,0),0),IFERROR(MATCH(%s,T_TIERRA_H,0),1)))' % MAT)
rx(35, "Electrodo 250.66", '=IF(BX31="","",IFERROR(IF(%s,INDEX(T_ELEC_AL,MATCH(BX31,T_ELEC_CALAL,0)),INDEX(T_ELEC_CU,MATCH(BX31,T_ELEC_CALCU,0))),""))' % ESAL)
rx(36, "Tubería (mm)", '=IF(BX31="","",IF(OR($F$3="CU-MC",$F$3="AL-MC"),"MC",IFERROR(VLOOKUP(BX31,T_CONDUIT,IF($F$4="RHW",3,IF($F$4="XHHW",4,2)),FALSE),"")))')
rx(37, "FAC tabla", '=IF(BX31="","",IFERROR(INDEX(CHOOSE(IF(%s,2,1),T_FAC_CU,T_FAC_AL),IF(BX15=1,1,IF(BX15>0.949,2,IF(BX15>0.899,3,IF(BX15>0.799,4,99)))),MATCH(BX31,T_FAC_H,0)),""))' % ESAL)
rx(38, "FAC ajustado", '=IF(BX37="","",BX37*IF(AND($C$8=1,OR($C$7=120,$C$7=230,$C$7=277)),1.1547,IF(OR(AND($C$8=1,OR($C$7=208,$C$7=240)),AND($C$8=2,$C$7=240)),1.15,1)))')
rx(39, "Corriente para caída (A)", '=IF(BX4>0,BX19,IF(BX5>0,BX20,BX21))'); rx(40, "Longitud (m)", '=N($C$12)')
rx(41, "Carga del transformador", '=IF($I$3="SI",IFERROR(BX14/$I$4,0),"")')
rx(42, "Caída en el transformador (%)", '=IF($I$3="SI",IFERROR(BX41*($I$5/SQRT(1+IF(N($I$6)>0,$I$6,4)^2))*(BX15+IF(N($I$6)>0,$I$6,4)*SQRT(1-BX15^2)),0),0)')
rx(43, "Tablero que lo alimenta", '=$C$10')
rx(44, "Voltaje del padre", '=IF($C$10="","",N(BX105))')
rx(45, "Bornes del padre (V)", '=IF($C$10="","",N(BX106))')
rx(46, "Icc del padre (A)", '=IF($C$10="","",N(BX107))')
rx(47, "Voltaje de partida (V)", '=IF($I$10="SI",$C$7,IF($I$3="SI",$C$7*(1-BX42/100),IF(AND($C$10<>"",N(BX44)=$C$7,N(BX45)>0),BX45,$C$7)))')
rx(48, "Caída en el alimentador (V)", '=IF(OR(BX38="",BX10=0),0,BX40*3.28*BX39*BX38/(10000*BX30))')
rx(49, "Voltaje en bornes (V)", '=BX47-BX48'); rx(50, "Caída acumulada (V)", '=$C$7-BX49'); rx(51, "Caída acumulada (%)", '=BX50*100/$C$7')
rx(52, "Voltaje L-N (V)", '=IF($C$7>490,$C$7/1.732,IF($C$7>240,277,IF($C$7=220,127,120)))')
rx(53, "Caída L-N (V)", '=BX48/IF($C$7=240,2,1.732)'); rx(54, "Bornes L-N (V)", '=BX52-($C$7-BX47)/IF($C$7=240,2,1.732)-BX53')
P2PARGS = 'BX31,%s,%s,BX30,BX40,$C$7,$C$8,ICC_LMAX' % (AIS, TUB)
rx(57, "Icc secundario del transformador (A)", '=IF($I$3<>"SI","",IFERROR(IF(IF($C$10<>"",N(BX46),N(ICC_RED)*1000)>0,(IF(N($I$7)>0,$I$7,IF(N(BX44)>0,BX44,$C$7))/$C$7)*IF($C$10<>"",N(BX46),N(ICC_RED)*1000)'
             '/(1+IF($C$8=3,1.732,1)*IF($C$10<>"",N(BX46),N(ICC_RED)*1000)*IF(N($I$7)>0,$I$7,IF(N(BX44)>0,BX44,$C$7))*$I$5/(100000*$I$4)),$I$4*1000/(IF($C$8=3,1.732,1)*$C$7*$I$5/100)),""))')
rx(55, "Icc por la alimentación normal (A)", '=IF(N($C$14)>0,$C$14*1000,IF($I$3="SI",IF(N(BX57)>0,P2P(BX57,%s),""),IF(AND($C$10<>"",N(BX44)=$C$7,N(BX46)>0),P2P(BX46,%s),"")))' % (P2PARGS, P2PARGS))
rx(59, "Icc de la UPS (A)", '=IF($I$10="SI",IF(N($I$12)>0,$I$12,2)*N($I$11)*1000/(IF($C$8=3,1.732,1)*$C$7),"")')
rx(56, "Icc en bornes (A)", '=IF(MAX(N(BX88),IF($I$10="SI",P2P(IF(N(BX55)>0,MIN(BX59,BX55),BX59),%s),N(BX55)))=0,"",MAX(N(BX88),IF($I$10="SI",P2P(IF(N(BX55)>0,MIN(BX59,BX55),BX59),%s),N(BX55))))' % (P2PARGS, P2PARGS))
rx(58, "Icc en bornes (kA)", '=IF(N(BX56)>0,BX56/1000,0)')
rx(60, "Marca", '=IF($C$15<>"",$C$15,MARCA_DEF)')
rx(61, "Espacios usados", '=MAX(SUMPRODUCT(MAX(($X$25:$X$124>0)*$B$25:$B$124)),MAX(0,$CA$25:$CB$124))')
rx(62, "Tablero de catálogo (tipo)", '=IF(N($F$17)>0,$F$17,TABLERO_SEL(BX60,$C$7,$C$5,N(BX24),BX61,RESERVA_ESP,$C$20,$C$16,N($C$17)))')
for i, (fila, campo) in enumerate([(63, "modelo"), (64, "barraF"), (65, "barraN"), (66, "barraT"), (67, "espacios"), (68, "familia"), (69, "ramales"), (70, "principales")]):
    rx(fila, "Tablero: " + campo, '=CAT_DATO(BX62,"%s")' % campo)
POLMAIN = 'IF($C$8=3,3,IF(AND($C$7>=208,$C$7<=240),2,$C$8))'
rx(71, "Interruptor principal (tipo)", '=IF(OR($C$16="Zapatas",BX24=""),"",IF(N($F$16)>0,$F$16,BREAKER_SEL(BX60,%s,BX24,$C$7,BX58,BX70,IF(OR($F$14="",$F$14="TM"),"STD",$F$14))))' % POLMAIN)
rx(72, "Principal: modelo", '=IF(BX71="",IF($C$16="Zapatas","ZAPATAS",""),BREAKER_REF(BX60,BX71,%s,BX24,$F$15="SI"))' % POLMAIN)
for fila, campo in [(73, "unidad"), (74, "sccr"), (75, "marco"), (76, "polos")]:
    rx(fila, "Principal: " + campo, '=IF(BX71="","",BREAKER_DATO(BX60,BX71,"%s"))' % campo)
rx(77, "SPD kA recomendados", '=IF($C$10="",250,IF(N(BX64)>=600,120,50))')
rx(78, "SPD (tipo)", '=IF(N($F$18)>0,$F$18,SPD_SEL(BX60,$C$6,$C$8,BX77))')
for fila, campo in [(79, "modelo"), (80, "kaLL"), (81, "kaLN"), (82, "montaje")]:
    rx(fila, "SPD: " + campo, '=IF(BX78="","",SPD_DATO(BX60,BX78,"%s",$C$6))' % campo)
rx(84, "Segunda acometida", '=$I$14="SI"')
rx(85, "2ª acometida: validación", '=IF(OR(NOT(BX84),BX24=""),"",CAL_AJUSTE(VLOOKUP(BX24*BX26/BX25,CHOOSE(IF(%s,2,1),T_AMPCU,T_AMPAL),2,TRUE),%s,BX25,%s,IF($F$11="","26-30",$F$11),IF(N($F$12)>0,$F$12,75),BX28,BX23,BX24,AUTO_AMP="SI",BX26))' % (ESAL, MAT, AIS))
rx(86, "2ª acometida: paralelos", '=IF(BX85="","",PARTEN(BX85,1))'); rx(87, "2ª acometida: calibre", '=IF(BX85="","",PARTE(BX85,2))')
ALTP2P = 'BX87,%s,%s,BX86,N($I$20),$C$7,$C$8,ICC_LMAX' % (AIS, TUB)
rx(88, "2ª acometida: Icc (A)", '=IF(OR(NOT(BX84),BX87=""),"",IF($I$16="Generador",IF(N($I$18)>0,P2P($I$18*1000/(IF($C$8=3,1.732,1)*$C$7*IF(N($I$19)>0,$I$19,12)/100),%s),""),IF($I$17<>"",IF(N(BX109)>0,P2P(BX109,%s),""),"")))' % (ALTP2P, ALTP2P))
rx(90, "2ª acometida: FAC ajustado", '=IF(BX87="","",IFERROR(INDEX(CHOOSE(IF(%s,2,1),T_FAC_CU,T_FAC_AL),IF(BX15=1,1,IF(BX15>0.949,2,IF(BX15>0.899,3,4))),MATCH(BX87,T_FAC_H,0))*IF(AND($C$8=2,$C$7=240),1.15,1),""))' % ESAL)
rx(89, "2ª acometida: caída (V)", '=IF(OR(BX87="",BX90=""),"",N($I$20)*3.28*BX39*BX90/(10000*BX86))')
rx(91, "2ª acometida: voltaje de partida", '=IF($I$16="Generador",$C$7,IF(N(BX110)>0,BX110,$C$7))')
rx(92, "2ª acometida: bornes (V)", '=IF(BX89="","",BX91-BX89)'); rx(93, "2ª acometida: caída (%)", '=IF(BX92="","",($C$7-BX92)*100/$C$7)')
rx(94, "Carga del generador", '=IF(AND(BX84,$I$16="Generador",N($I$18)>0),BX14/$I$18,"")')
rx(95, "2ª acometida: conductores", '=IF(BX87="","",IF(BX86=1,"",BX86&"x")&IF($C$8=3,"3#","2#")&BX87)')
rx(98, "Alimentado desde", '=IF($C$10<>"",IF(BX108<>"",BX108,$C$10),$C$11&"")')
rx(99, "Fases (texto)", '=IF(BX31="","",BX32&BX31)'); rx(100, "Neutro (texto)", '=IF(BX33="","",BX104&BX33)'); rx(101, "Tierra (texto)", '=IF(BX34="","",BX104&BX34)')
rx(102, "Tubería (texto)", '=IF(BX36="","",BX104&BX36)'); rx(103, "Fuente del Icc", '=IF(N($C$14)>0,"manual",IF($I$3="SI","transformador",IF($I$10="SI","UPS",IF(N(BX55)>0,"cascada",""))))')
for fila_, et_ in [(105, "Enlace: voltaje del padre"), (106, "Enlace: bornes del padre"), (107, "Enlace: Icc del padre"), (108, "Enlace: nombre del padre"), (109, "Enlace: Icc del origen del bypass"), (110, "Enlace: bornes del origen del bypass")]:
    rx(fila_, et_, '=""')     # las macros escriben aquí la referencia directa al tablero padre / de origen
rx(104, "Prefijo paralelos", '=IF(BX30=1,"",BX30&"x")')
rx(97, "Avisos", '=TRIM(IF(BX51>CVMAX_ALIM,"Caída acumulada "&FIXED(BX51,2)&" % > "&CVMAX_ALIM&" %. ","")&IF(AND(BX29<>"",PARTE(BX29,3)="0"),"Alimentador no cumple ampacidad corregida. ","")'
             '&IF(BX62="","Sin tablero de catálogo. ","")&IF(AND(N(BX67)>0,N(BX67)<BX61),"Faltan espacios. ","")&IF(AND(BX74<>"",N(BX74)<BX58),"SCCR del principal < Icc. ","")&IF(BX7>DESB_MAX,"Desbalance "&FIXED(BX7,1)&" %. ","")'
             '&IF(AND(BX41<>"",N(BX41)>1),"Transformador sobrecargado. ","")&IF(AND(N(BX44)>0,N(BX44)<>$C$7,$I$3<>"SI"),"Voltaje distinto al del padre: agregue transformador. ","")'
             '&IF(AND(BX93<>"",N(BX93)>CVMAX_ALIM),"Caída de la 2ª acometida alta. ","")&IF(AND($I$14="SI",OR($I$15="ATS",$I$15="MTS"),N($I$21)>0,N($I$21)<BX58),"SCCR del "&$I$15&" < Icc. ","")'
             '&IF(BX80<>"",IF(N(BX80)<BX77,"SPD de "&BX80&" kA < "&BX77&" kA recomendados (242). ",""),"")&IF(N(BX24)>=1200,"240.87: reducción de energía de arco. ","")'
             '&IF(COUNTIF($BU$25:$BU$124,"?*")>0,COUNTIF($BU$25:$BU$124,"?*")&" circuito(s) con aviso. ",""))')
mc.column_dimensions["CA"].hidden = True; mc.column_dimensions["CB"].hidden = True
mc.column_dimensions["BW"].width = 40; mc.column_dimensions["BX"].width = 22

# panel visible de resultados
panel = [("K3", "kVA conectados", "BX10"), ("K4", "kVA demandados", "BX14"), ("K5", "Factor de demanda", "BX12"), ("K6", "Corriente A / B / C (A)", None), ("K7", "Protección principal (A)", "BX24"),
         ("K8", "Conductores", None), ("K9", "Tubería", "BX102"), ("K10", "Voltaje de partida / bornes (V)", None), ("K11", "Caída acumulada (%)", "BX51"), ("K12", "Icc en bornes (kA)", "BX58"),
         ("K13", "Fuente del Icc", "BX103"), ("K14", "Tablero (catálogo)", None), ("K15", "Interruptor principal", None), ("K16", "Supresor (SPD)", None), ("K17", "Desbalance (%)", "BX7"),
         ("K18", "Validación del alimentador", None), ("K19", "Segunda acometida", None), ("K20", "AVISOS", "BX97")]
for cel, et, ref in panel:
    f(mc, cel, et, size=9, bold=True, fill=AZUL, borde=True)
    v = "=" + ref if ref else None
    r = cel[1:]
    if cel == "K6": v = '=FIXED(BX19,1)&" / "&FIXED(BX20,1)&" / "&FIXED(BX21,1)'
    if cel == "K8": v = '=BX99&IF(BX100<>""," + "&BX100&" N","")&IF(BX101<>""," + "&BX101&" T","")&" "&IF($F$3="","CU",$F$3)&" "&IF($F$4="","XHHW-2",$F$4)'
    if cel == "K10": v = '=FIXED(BX47,2)&" / "&FIXED(BX49,2)'
    if cel == "K14": v = '=BX63&" ("&BX60&", "&BX68&") · "&BX64&" A · "&BX67&" esp."'
    if cel == "K15": v = '=BX72&" "&BX73&" · "&BX74&" kA"'
    if cel == "K16": v = '=BX79&" · "&BX80&"/"&BX81&" kA"'
    if cel == "K18": v = '=IF(BX29="","",IF(PARTE(BX29,3)="1","Cumple","NO CUMPLE")&" · disponible "&FIXED(PARTEN(BX29,4),1)&" A vs "&FIXED(N(BX23),1)&" A")'
    if cel == "K19": v = '=IF(BX84,$I$15&" · "&BX95&" · ΔV "&FIXED(N(BX93),2)&" % · Icc "&FIXED(N(BX88)/1000,2)&" kA","")'
    c = mc.cell(int(r), CI("L")); c.value = v; c.font = Font(name=FN, size=10, bold=True); c.border = BORDE
    mc.merge_cells("L%s:R%s" % (r, r))
mc.column_dimensions["K"].width = 30
for c, fm in [("L11", "0.00"), ("L12", "0.00"), ("L3", "0.00"), ("L4", "0.00"), ("L5", "0.000"), ("L17", "0.00")]: mc[c].number_format = fm

# circuitos (filas 25 a 124: primero las posiciones impares 1..99 y luego las pares 2..100, como el Excel)
cab = ["Fijo", "Pos.", "Polo 2", "Polo 3", "Detalle", "Descripción", "kVA", "Long. (m)", "Prot. %", "F. uso", "F. div.", "Material", "Aislam.", "# par.", "Aum. %", "Prot. (A)", "Unidad", "Breaker (tipo)",
       "Tablero derivado", "Bypass de", "Circuito Revit", "Origen",
       "Descripción", "kVA", "V", "Fases", "Hilos", "F. uso", "F. div.", "FP", "Tipo carga", "Mult.", "fA", "fB", "fC", "kVA A", "kVA B", "kVA C", "dem A", "dem B", "dem C", "I A", "I B", "I C",
       "Amp. req.", "Prot. (A)", "# par.", "Aum.", "Calibre", "Prefijo", "Neutro", "Tierra", "Tubo", "FAC", "FAC aj.", "V real", "ΔV (V)", "V-ΔV", "ΔV total (V)", "ΔV total (%)",
       "Polos bk", "Unidad", "Breaker (tipo)", "Modelo de referencia", "SCCR (kA)", "Serie (kA)", "Unidad real", "Marco", "Fases", "Neutro", "Tierra", "Tubo", "Aviso"]
f(mc, "A22", "CIRCUITOS RAMALES — celdas verdes: datos; vacías = valor automático. Filas: posiciones impares (1–99) y luego pares (2–100).", bold=True, size=10)
for j, t in enumerate(cab):
    c = mc.cell(24, 1 + j, t); c.font = Font(name=FN, bold=True, size=9); c.fill = GRIS if j < 22 else GRIS2; c.alignment = CENT; c.border = BORDE
mc["A23"] = "ENTRADAS"; mc["W23"] = "CÁLCULO (no editar)"
mc["A23"].font = Font(name=FN, bold=True, color="006600"); mc["W23"].font = Font(name=FN, bold=True, color="7F7F7F")
for i in range(100):
    r = R0 + i
    pos = 2 * i + 1 if i < 50 else 2 * (i - 50) + 2
    mc.cell(r, 2, pos).font = Font(name=FN, bold=True)
    for col in list("ACDEFGHIJKLMNOPQRSTUV"):
        c = mc[col + str(r)]; c.fill = VERDE; c.border = BORDE; c.font = Font(name=FN, size=9)
    def cur(u, x):
        return ('IF({u}=0,0,IF(AND($Y{r}=120,$Z{r}=1),{x}*1000/$Y{r},IF(AND($Y{r}=208,$Z{r}=1),{x}*1000/$Y{r},IF(AND($Y{r}=208,$Z{r}=2),{u}*2*1000/$Y{r},'
                'IF(AND($Y{r}=230,$Z{r}=1),{x}*1000/$Y{r},IF(AND($Y{r}=240,$Z{r}=1),{x}*1000/($Y{r}/2),IF(AND($Y{r}=277,$Z{r}=1),{x}*1000/$Y{r},{x}*1000/($Y{r}/1.732))))))))').format(u=u, x=x, r=r)
    MATR = 'IF($L{r}="","CU",$L{r})'.format(r=r)
    ESALR = 'OR($L{r}="AL",$L{r}="AL-MC")'.format(r=r)
    forms = {
        "W": '=IF($F{r}<>"",$F{r},IFERROR(VLOOKUP($E{r},T_DET,3,FALSE),""))', "X": '=IF(ISNUMBER($G{r}),$G{r},0)',
        "Y": '=IFERROR(VLOOKUP($E{r},T_DET,4,FALSE),0)', "Z": '=IFERROR(VLOOKUP($E{r},T_DET,5,FALSE),0)', "AA": '=IFERROR(VLOOKUP($E{r},T_DET,6,FALSE),0)',
        "AB": '=IF(ISNUMBER($J{r}),$J{r},IFERROR(VLOOKUP($E{r},T_DET,7,FALSE),1))', "AC": '=IF(N($K{r})>0,$K{r},IFERROR(VLOOKUP($E{r},T_DET,8,FALSE),1))',
        "AD": '=IFERROR(VLOOKUP($E{r},T_DET,9,FALSE),1)', "AE": '=IFERROR(VLOOKUP($E{r},T_DET,2,FALSE),"")',
        "AF": '=IF(ISNUMBER($I{r}),$I{r}/100,IFERROR(VLOOKUP($E{r},T_DET,11,FALSE),1.25))',
        "AG": '=IF(N($B{r})>0,MOD(INT(($B{r}-1)/2),IF($C$8=3,3,2)),-1)', "AH": '=IF(N($CA{r})>0,MOD(INT(($CA{r}-1)/2),IF($C$8=3,3,2)),-1)', "AI": '=IF(N($CB{r})>0,MOD(INT(($CB{r}-1)/2),IF($C$8=3,3,2)),-1)',
        "CA": '=IF(N($C{r})>0,$C{r},"")', "CB": '=IF(N($D{r})>0,$D{r},"")',
        "AJ": '=IF(OR($X{r}=0,$Z{r}=0),0,IF($C$8=1,$X{r},IF(OR($AG{r}=0,$AH{r}=0,$AI{r}=0),$X{r}/$Z{r},0)))',
        "AK": '=IF(OR($X{r}=0,$Z{r}=0,$C$8=1),0,IF(OR($AG{r}=1,$AH{r}=1,$AI{r}=1),$X{r}/$Z{r},0))',
        "AL": '=IF(OR($X{r}=0,$Z{r}=0,$C$8<3),0,IF(OR($AG{r}=2,$AH{r}=2,$AI{r}=2),$X{r}/$Z{r},0))',
        "AM": '=$AJ{r}*$AB{r}/$AC{r}', "AN": '=$AK{r}*$AB{r}/$AC{r}', "AO": '=$AL{r}*$AB{r}/$AC{r}',
        "AP": "=" + cur("$AJ{r}".format(r=r), "$AM{r}".format(r=r)), "AQ": "=" + cur("$AK{r}".format(r=r), "$AN{r}".format(r=r)), "AR": "=" + cur("$AL{r}".format(r=r), "$AO{r}".format(r=r)),
        "AS": '=IF($X{r}=0,"",IF($C$8=3,MAX($AP{r}:$AR{r}),IF($AJ{r}>0,$AP{r},IF($AK{r}>0,$AQ{r},$AR{r})))*$AF{r})',
        "AT": '=IF($AS{r}="","",IF(N($P{r})>0,$P{r},VLOOKUP($AS{r},T_PROT,2,TRUE)))',
        "AU": '=IF($AS{r}="",1,IF(N($N{r})>0,$N{r},IF($AS{r}>1600,5,IF($AS{r}>1000,4,IF($AS{r}>600,3,IF($AS{r}>300,2,1))))))',
        "AV": '=1+N($O{r})/100',
        "AW": '=IF($AT{r}="","",VLOOKUP($AT{r}*$AV{r}/$AU{r},CHOOSE(IF(' + ESALR + ',2,1),T_AMPCU,T_AMPAL),2,TRUE))',
        "AX": '=IF($AW{r}="","",IF($AU{r}=1,"",$AU{r}&"x")&IF($Z{r}=3,"3#",IF(AND($Y{r}>=208,$Y{r}<=240),"2#","")))',
        "AY": '=IF($AW{r}="","",IF($AA{r}=0,"",IF(AND($Y{r}>=208,$Z{r}=3,$AA{r}=4),$AW{r},IF(AND($Z{r}=3,$AA{r}=3),"",IF(AND($Y{r}>=208,$Z{r}=1,$AA{r}=3),"",$AW{r})))))',
        "AZ": '=IF($AT{r}="","",INDEX(T_TIERRA,MATCH(TRUE,INDEX(T_TIERRA_A>=$AT{r},0),0),IFERROR(MATCH(' + MATR + ',T_TIERRA_H,0),1)))',
        "BA": '=IF($AW{r}="","",IF(OR($L{r}="CU-MC",$L{r}="AL-MC"),"MC",IFERROR(VLOOKUP($AW{r},T_CONDUIT,IF($M{r}="RHW",3,IF($M{r}="XHHW",4,2)),FALSE),"")))',
        "BB": '=IF($AW{r}="","",IFERROR(INDEX(CHOOSE(IF(' + ESALR + ',2,1),T_FAC_CU,T_FAC_AL),IF($AD{r}=1,1,IF($AD{r}>0.949,2,IF($AD{r}>0.899,3,IF($AD{r}>0.799,4,99)))),MATCH($AW{r},T_FAC_H,0)),""))',
        "BC": '=IF($BB{r}="","",$BB{r}*IF(AND($Z{r}=1,OR($Y{r}=120,$Y{r}=230,$Y{r}=277)),1.1547,IF(AND($Z{r}=1,OR($Y{r}=208,$Y{r}=240)),1.15,1)))',
        "BD": '=IF($X{r}=0,"",IF($Y{r}=$C$7,$BX$49,$BX$54))',
        "BE": '=IF(OR($X{r}=0,$BC{r}=""),"",N($H{r})*3.28*IF($AJ{r}>0,$AP{r},IF($AK{r}>0,$AQ{r},$AR{r}))*$BC{r}/(10000*$AU{r}))',
        "BF": '=IF($BE{r}="","",$BD{r}-$BE{r})', "BG": '=IF($BF{r}="","",$Y{r}-$BF{r})', "BH": '=IF($BG{r}="","",$BG{r}*100/$Y{r})',
        "BI": '=IF($X{r}=0,"",IF($Z{r}=3,3,IF(AND($Y{r}>=208,$Y{r}<=240),2,MAX($Z{r},1))))',
        "BJ": '=IF($X{r}=0,"",IF($Q{r}<>"",$Q{r},NEC_UNIDAD($W{r},$E{r},$Y{r},$Z{r},N($AT{r}),$C$7,OCUPACION)))',
        "BK": '=IF($X{r}=0,"",IF(N($R{r})>0,$R{r},BREAKER_SEL($BX$60,$BI{r},$AT{r},$C$7,$BX$58,$BX$69,$BJ{r})))',
        "BL": '=IF($BK{r}="","",BREAKER_REF($BX$60,$BK{r},$BI{r},$AT{r},FALSE))', "BM": '=IF($BK{r}="","",BREAKER_DATO($BX$60,$BK{r},"sccr"))',
        "BN": '=IF(OR($BK{r}="",N($BM{r})>=$BX$58,$BX$71="",N($BX$74)<$BX$58),"",SERIE_KA($BX$60,BREAKER_DATO($BX$60,$BK{r},"modelo"),$BX$72,N($BX$24),$C$7,$BX$58))',
        "BO": '=IF($BK{r}="","",BREAKER_DATO($BX$60,$BK{r},"unidad"))', "BP": '=IF($BK{r}="","",BREAKER_DATO($BX$60,$BK{r},"marco"))',
        "BQ": '=IF($AW{r}="","",$AX{r}&$AW{r})', "BR": '=IF(OR($AY{r}="",$AW{r}=""),"",IF($AU{r}=1,"",$AU{r}&"x")&$AY{r})',
        "BS": '=IF($AZ{r}="","",IF($AU{r}=1,"",$AU{r}&"x")&$AZ{r})', "BT": '=IF($BA{r}="","",IF($AU{r}=1,"",$AU{r}&"x")&$BA{r})',
        "BU": '=IF($X{r}=0,"",TRIM(IF($E{r}="","Sin detalle de carga. ","")&IF(N($BH{r})>CVMAX_TOT,"Caída "&FIXED($BH{r},2)&" %. ","")&IF(AND($BM{r}<>"",N($BM{r})<$BX$58,$BN{r}=""),"SCCR "&$BM{r}&" kA < Icc. ","")&IF(AND($Z{r}>0,COUNT($B{r}:$D{r})>$Z{r}),"Más polos que fases del detalle. ","")&IF(AND(COUNT($B{r}:$D{r})<$Z{r},$Z{r}>1),"Faltan polos: el detalle es de "&$Z{r}&" fases. ","")&IF(COUNTIF($CA$25:$CB$124,$B{r})>0,"Posición ocupada por otro circuito. ","")&IF($BN{r}<>"","Serie "&$BN{r}&" kA (110.22(C)). ","")))',
    }
    for col, fm in forms.items():
        c = mc[col + str(r)]; c.value = fm.format(r=r); c.font = Font(name=FN, size=9)
for col, fm in [("G", "0.00"), ("X", "0.00"), ("AJ", "0.00"), ("AK", "0.00"), ("AL", "0.00"), ("AP", "0.0"), ("AQ", "0.0"), ("AR", "0.0"), ("AS", "0.0"), ("BE", "0.00"), ("BH", "0.00")]:
    for r in range(R0, RN + 1): mc[col + str(r)].number_format = fm
mc.column_dimensions["F"].width = 34; mc.column_dimensions["W"].width = 30; mc.column_dimensions["BL"].width = 22; mc.column_dimensions["BU"].width = 40
for col in ["E", "S", "T"]:
    pass
dv = DataValidation(type="list", formula1="=L_MATERIALES", allow_blank=True); mc.add_data_validation(dv); dv.add("L25:L124")
dv = DataValidation(type="list", formula1="=L_AISLAMIENTOS", allow_blank=True); mc.add_data_validation(dv); dv.add("M25:M124")
dv = DataValidation(type="list", formula1="=L_UNIDADES", allow_blank=True); mc.add_data_validation(dv); dv.add("Q25:Q124")
dv = DataValidation(type="list", formula1="=L_DET", allow_blank=True); mc.add_data_validation(dv); dv.add("E25:E124")
dv = DataValidation(type="list", formula1='"SI"', allow_blank=True); mc.add_data_validation(dv); dv.add("A25:A124")
mc.freeze_panes = "C25"

# factores de demanda (filas 127 a 146)
f(mc, "B126", "FACTORES DE DEMANDA Y DIVERSIDAD (deje F. demanda vacío para usar el del catálogo; tomas: 100 % de 10 kVA + 50 % del resto, 220.44)", bold=True)
for j, t in enumerate(["Código", "Tipo de carga", "Método", "Circuitos", "kVA conectados", "Reserva", "kVA totales", "F. demanda (manual)", "F. demanda", "F. diversidad", "kVA demandados"]):
    c = mc.cell(127, 2 + j, t); c.font = Font(name=FN, bold=True, size=9); c.fill = GRIS; c.alignment = CENT; c.border = BORDE
EXCL = 'IF(CARGA_DERIV="demandada",COUNTIFS($S$25:$S$124,"",{0}),{1})'
for k in range(16):
    r = 128 + k
    mc["B%d" % r] = '=IFERROR(IF(INDEX(T_TIPOS,%d,1)="","",INDEX(T_TIPOS,%d,1)),"")' % (k + 1, k + 1)
    mc["C%d" % r] = '=IF($B%d="","",INDEX(T_TIPOS,%d,2))' % (r, k + 1)
    mc["D%d" % r] = '=IF($B%d="","",INDEX(T_TIPOS,%d,3))' % (r, k + 1)
    mc["E%d" % r] = '=IF($B{r}="",0,IF(CARGA_DERIV="demandada",COUNTIFS($AE$25:$AE$124,$B{r},$X$25:$X$124,">0",$T$25:$T$124,"",$S$25:$S$124,""),COUNTIFS($AE$25:$AE$124,$B{r},$X$25:$X$124,">0",$T$25:$T$124,"")))'.format(r=r)
    mc["F%d" % r] = '=IF($B{r}="",0,IF(CARGA_DERIV="demandada",SUMIFS($X$25:$X$124,$AE$25:$AE$124,$B{r},$T$25:$T$124,"",$S$25:$S$124,""),SUMIFS($X$25:$X$124,$AE$25:$AE$124,$B{r},$T$25:$T$124,"")))'.format(r=r)
    mc["G%d" % r] = '=F{r}*N($C$13)'.format(r=r); mc["H%d" % r] = '=F{r}+G{r}'.format(r=r)
    mc["J%d" % r] = '=IF($B{r}="",1,IF(ISNUMBER($I{r}),$I{r},IF($D{r}="tomas",IF($H{r}>0,IF($H{r}>10,($H{r}-10)*0.5+10,$H{r})/$H{r},1),IF($D{r}="cocina",LOOKUP($E{r},{{0,1,2,3,4,5,6}},{{1,1,1,0.9,0.8,0.7,0.65}}),IF($D{r}="220.53",IF($E{r}>=4,0.75,1),N(INDEX(T_TIPOS,{k},4)))))))'.format(r=r, k=k + 1)
    mc["K%d" % r] = '=IF($B{r}="",1,IF(N(INDEX(T_TIPOS,{k},5))>0,INDEX(T_TIPOS,{k},5),1))'.format(r=r, k=k + 1)
    mc["L%d" % r] = '=IF(AND($D{r}="tomas",NOT(ISNUMBER($I{r}))),IF($H{r}>10,($H{r}-10)*0.5+10,$H{r}),$H{r}*$J{r}/$K{r})'.format(r=r)
    mc["I%d" % r].fill = VERDE
r = 144
mc["C%d" % r] = "TABLEROS DERIVADOS (demanda ya aplicada)"
mc["E%d" % r] = '=IF(CARGA_DERIV="demandada",COUNTIFS($S$25:$S$124,"<>",$X$25:$X$124,">0",$T$25:$T$124,""),0)'
mc["F%d" % r] = '=IF(CARGA_DERIV="demandada",SUMIFS($X$25:$X$124,$S$25:$S$124,"<>",$T$25:$T$124,""),0)'
mc["G%d" % r] = '=F144*N($C$13)'; mc["H%d" % r] = '=F144+G144'; mc["J%d" % r] = 1; mc["K%d" % r] = 1; mc["L%d" % r] = '=H144'
r = 146
mc["C146"] = "TOTALES"; mc["C146"].font = Font(name=FN, bold=True)
for col in "FGHL": mc["%s146" % col] = '=SUM(%s128:%s144)' % (col, col); mc["%s146" % col].font = Font(name=FN, bold=True)
mc["J146"] = '=BX12'
for rr in range(128, 147):
    for col in "FGHJL": mc["%s%d" % (col, rr)].number_format = "0.00"

# ======================= VISTA_MACHOTE: cuadro de cargas con el diseño de la hoja TABLEROS 3F =======================
vs = wb.create_sheet("VISTA_MACHOTE", 1)
anchos = {"A": 1.5, "B": 22.8, "C": 2.8, "D": 6.2, "E": 5, "F": 4.8, "G": 21.1, "H": 32.1, "I": 18.1, "J": 15.8, "K": 13.8, "L": 13.9, "M": 13.5, "N": 16.1, "O": 22.2, "P": 13.9, "Q": 12, "R": 13, "S": 18.1, "T": 16.2, "U": 13.2, "V": 15.8, "W": 15.2,
          "X": 18.1, "Y": 2.8, "Z": 6.1, "AA": 5.1, "AB": 5.5, "AC": 7, "AD": 7.5, "AE": 6, "AF": 5.8, "AG": 7.5, "AH": 5.8, "AI": 12.5, "AJ": 4.5}
for k, w in anchos.items(): vs.column_dimensions[k].width = w
M = "MC_MACHOTE!"
EST = {"tit": (48, True, "006600", None), "alimL": (18, True, None, GRIS), "alimV": (18, True, None, None), "barra": (18, True, None, GRIS), "th": (12, True, None, None),
       "td": (12, True, None, None), "tdr": (12, True, None, None), "gris": (12, True, None, GRIS), "tot": (12, True, None, GRIS), "ley": (9, False, None, None)}
def V_(y, a, b, val, est, fmt=None, y2=None):
    y2 = y2 or y
    sz, bo, col, fill = EST[est]
    for yy in range(y, y2 + 1):
        for x in range(CI(a), CI(b) + 1):
            c = vs.cell(yy, x)
            if est != "ley": c.border = BORDEM if est in ("tit", "alimV") else BORDE
            if fill: c.fill = fill
    c = vs["%s%d" % (a, y)]; c.value = val
    c.font = Font(name=FN, size=sz, bold=bo, color=col)
    c.alignment = Alignment(horizontal="right" if est == "tdr" else ("left" if est == "ley" else "center"), vertical="center" if est != "ley" else "top", wrap_text=est in ("th", "alimV", "ley", "tot"))
    if fmt: c.number_format = fmt
    if (a != b) or (y2 != y): vs.merge_cells("%s%d:%s%d" % (a, y, b, y2))
X = lambda ref: "=" + M + ref
V_(1, "B", "L", X("$BX$2"), "tit", y2=2); V_(1, "M", "O", "Alimentado desde:", "alimL"); V_(2, "M", "O", X("$BX$98"), "alimV"); V_(1, "P", "AJ", None, "alimV", y2=2)
vs.row_dimensions[1].height = 61; vs.row_dimensions[2].height = 56
V_(3, "B", "O", "Datos de cálculos eléctricos", "barra"); V_(3, "P", "AJ", "Especificación del tablero", "barra"); vs.row_dimensions[3].height = 48
for a, b, t in [("B", "F", "KVA Conectados"), ("G", "G", "KVA Demandados"), ("H", "H", "kVA Reserva"), ("I", "I", "Factor demanda"), ("J", "J", "Factor diversidad"), ("K", "K", "Factor   potencia"),
                ("O", "O", "Icc disponible (kA)"), ("P", "Q", "Tensión Nominal (Voltios)"), ("R", "R", "Fases"), ("S", "S", "Hilos"), ("T", "U", "Modelo de referencia"), ("V", "V", "Fabricante"),
                ("W", "W", "Montaje"), ("AF", "AG", "Espacios"), ("AH", "AJ", "% DESBALANCEO MAXIMO")]:
    V_(4, a, b, t, "th", y2=5)
V_(4, "L", "N", "Amperios totales por fase", "th"); V_(5, "L", "L", "Fase A", "th"); V_(5, "M", "M", "Fase B", "th"); V_(5, "N", "N", "Fase C", "th")
V_(4, "X", "AE", "Amperaje de Barras", "th"); V_(5, "X", "AA", "Fases", "th"); V_(5, "AB", "AC", "Neutro", "th"); V_(5, "AD", "AE", "Tierra", "th")
for a, b, ref, fm in [("B", "F", "$BX$10", "0.00"), ("G", "G", "$BX$14", "0.00"), ("H", "H", "$C$13", "0%"), ("I", "I", "$BX$12", "0.00"), ("J", "J", "$BX$13", "0.00"), ("K", "K", "$BX$15", "0.00"),
                      ("L", "L", "$BX$19", "0.0"), ("M", "M", "$BX$20", "0.0"), ("N", "N", "$BX$21", "0.0"), ("O", "O", "$BX$58", "0.00"), ("P", "Q", "$C$6", None), ("R", "R", "$C$8", None), ("S", "S", "$C$9", None),
                      ("T", "U", "$BX$63", None), ("V", "V", "$BX$60", None), ("W", "W", "$C$18", None), ("X", "AA", "$BX$64", None), ("AB", "AC", "$BX$65", None), ("AD", "AE", "$BX$66", None), ("AF", "AG", "$BX$67", None), ("AH", "AJ", "$BX$7", "0.00")]:
    V_(6, a, b, X(ref), "td", fm)
for y in (4, 5): vs.row_dimensions[y].height = 23
vs.row_dimensions[6].height = 38
V_(7, "B", "O", "Datos del alimentador/acometida", "barra"); V_(7, "P", "W", "Datos de Interruptor Principal", "barra"); V_(7, "X", "AJ", "Datos Supresor (SPD)", "barra")
for a, b, t in [("B", "F", "Fases (AWG)"), ("G", "G", "Neutro (AWG)"), ("H", "H", "Tierra (AWG)"), ("I", "I", "Material"), ("J", "J", "Aislamiento"), ("K", "K", "Tubería (Ø mm)"), ("L", "L", "Longitud (m)"),
                ("M", "M", "Voltaje (V)"), ("N", "N", "Caída voltaje (V)"), ("O", "O", "Caída voltaje total (%)"), ("P", "R", "Modelo de referencia"), ("S", "S", "Amperaje (A)"), ("T", "T", "Marco (A)"),
                ("U", "U", "Tipo"), ("V", "V", "Polos"), ("W", "W", "SCCR (kA)"), ("X", "AC", "Modelo de referencia"), ("AD", "AF", "Capacidad de supresión (kA)"), ("AI", "AJ", "Montaje")]:
    V_(8, a, b, t, "th")
V_(8, "AG", "AH", None, "th", y2=9)
for a, b, ref, fm in [("B", "F", "$BX$99", None), ("G", "G", "$BX$100", None), ("H", "H", "$BX$101", None), ("I", "I", "$F$3", None), ("J", "J", "$F$4", None), ("K", "K", "$BX$102", None), ("L", "L", "$BX$40", "0.0"),
                      ("M", "M", "$BX$49", "0.00"), ("N", "N", "$BX$50", "0.00"), ("O", "O", "$BX$51", "0.00"), ("P", "R", "$BX$72", None), ("S", "S", "$BX$24", None), ("T", "T", "$BX$75", None), ("U", "U", "$BX$73", None),
                      ("V", "V", "$BX$76", None), ("W", "W", "$BX$74", None), ("X", "AC", "$BX$79", None), ("AD", "AF", "$BX$80", None), ("AI", "AJ", "$BX$82", None)]:
    V_(9, a, b, X(ref), "td", fm)
for y in (7, 8, 9): vs.row_dimensions[y].height = 38
V_(10, "B", "W", "Datos de circuitos ramales / alimentadores", "barra"); V_(10, "X", "AJ", "Posición en el tablero", "barra"); vs.row_dimensions[10].height = 38
for a, b, t in [("B", "F", "ID de circuito"), ("G", "H", "Descripción de carga"), ("I", "I", "kVA Conectados"), ("J", "J", "Caída voltaje total (%)"), ("X", "AB", "ID de circuito")]:
    V_(11, a, b, t, "th", y2=12)
V_(11, "K", "Q", "Datos de Interruptores Ramales", "th"); V_(11, "R", "V", "Datos conductores ramales / alimentadores", "th"); V_(11, "W", "W", "Tubo EMT", "th")
for a, b, t in [("K", "L", "Modelo de referencia"), ("M", "M", "Marco (A)"), ("N", "N", "Amperios (A)"), ("O", "O", "Tipo de unidad"), ("P", "P", "Polos"), ("Q", "Q", "SCCR (kA)"), ("R", "R", "Fases"),
                ("S", "S", "Neutro"), ("T", "T", "Tierra"), ("U", "U", "Tipo"), ("V", "V", "Aislamiento"), ("W", "W", "(Ø mm)")]:
    V_(12, a, b, t, "th")
FC = [("AC", "AE"), ("AF", "AH"), ("AI", "AJ")]
for i, (a, b) in enumerate(FC): V_(11, a, b, "KVA", "th"); V_(12, a, b, "FASE " + "ABC"[i], "th")
vs.row_dimensions[11].height = 22; vs.row_dimensions[12].height = 45
gris_vacio = PatternFill("solid", fgColor="D9D9D9")
y = 13
for i in range(100):
    p = 2 * i + 1 if i < 50 else 2 * (i - 50) + 2
    rr = R0 + i                                           # fila de esa posición en la memoria
    ph = ((p - 1) // 2) % 3
    hay = M + "$X$%d>0" % rr
    def S(expr, vacio='""'): return '=IF(%s,%s,%s)' % (hay, expr, vacio)
    V_(y, "B", "B", "=" + M + "$C$3", "tdr"); V_(y, "C", "C", "-", "td"); V_(y, "D", "D", p, "tdr")
    V_(y, "E", "E", S(M + "$CA$%d" % rr), "tdr"); V_(y, "F", "F", S(M + "$CB$%d" % rr), "tdr")
    V_(y, "G", "H", S(M + "$W$%d" % rr), "td"); V_(y, "I", "I", S(M + "$X$%d" % rr), "td", "0.00"); V_(y, "J", "J", S(M + "$BH$%d" % rr), "td", "0.00")
    V_(y, "K", "L", S(M + "$BL$%d" % rr), "td"); V_(y, "M", "M", S("IF(N(%s$BP$%d)>0,%s$BP$%d,\"\")" % (M, rr, M, rr)), "td"); V_(y, "N", "N", S(M + "$AT$%d" % rr), "td")
    V_(y, "O", "O", S(M + "$BO$%d" % rr), "td"); V_(y, "P", "P", S(M + "$BI$%d" % rr), "td"); V_(y, "Q", "Q", S(M + "$BM$%d" % rr), "td")
    V_(y, "R", "R", S(M + "$BQ$%d" % rr), "td"); V_(y, "S", "S", S(M + "$BR$%d" % rr), "td"); V_(y, "T", "T", S(M + "$BS$%d" % rr), "td")
    V_(y, "U", "U", S('IF(%s$L$%d="","CU",IF(%s$L$%d="AL-MC","AL",%s$L$%d))' % (M, rr, M, rr, M, rr)), "td"); V_(y, "V", "V", S('IF(%s$M$%d="","THHN",%s$M$%d)' % (M, rr, M, rr)), "td")
    V_(y, "W", "W", S(M + "$BT$%d" % rr), "td")
    V_(y, "X", "X", "=" + M + "$C$3", "tdr"); V_(y, "Y", "Y", "-", "td"); V_(y, "Z", "Z", p, "tdr"); V_(y, "AA", "AA", "=E%d" % y, "tdr"); V_(y, "AB", "AB", "=F%d" % y, "tdr")
    for k, (a, b) in enumerate(FC):
        V_(y, a, b, S("IF(%s$%s$%d>0,%s$%s$%d,\"\")" % (M, ["AJ", "AK", "AL"][k], rr, M, ["AJ", "AK", "AL"][k], rr)), "td", "0.00")
    vs["AL%d" % y] = '=COUNTIF(%s$CA$25:$CB$124,%d)>0' % (M, p)        # posición ocupada por un circuito de varios polos
    vs["AM%d" % y] = p
    vs.row_dimensions[y].height = 23.45
    y += 1
ult = y - 1
vs.conditional_formatting.add("G13:H%d" % ult, FormulaRule(formula=['$I13=""'], fill=gris_vacio))
for k, (a, b) in enumerate(FC):
    vs.conditional_formatting.add("%s13:%s%d" % (a, b, ult), FormulaRule(formula=['AND($%s13="",OR($AL13,MOD(INT(($AM13-1)/2),3)<>%d))' % (a, k)], fill=gris_vacio))
vs.column_dimensions["AL"].hidden = True; vs.column_dimensions["AM"].hidden = True
V_(y, "B", "W", "=" + M + '$BW$1', "ley", y2=y + 1)
vs["B%d" % y] = D.get("leyendaUnidades", "")
V_(y, "X", "AB", "kVA Conectados sin reserva", "tot", y2=y + 1)
for k, (a, b) in enumerate(FC): V_(y, a, b, "=" + M + "$BX$%d" % (4 + k), "tot", "0.00", y2=y + 1)
vs.sheet_view.showGridLines = False; vs.sheet_view.zoomScale = 60
vs.page_setup.orientation = "landscape"; vs.page_setup.fitToWidth = 1; vs.page_setup.fitToHeight = 0; vs.sheet_properties.pageSetUpPr.fitToPage = True
vs.print_area = "A1:AJ%d" % (y + 1)
try:
    img = XLImage(LOGO); img.width, img.height = 300, 107; vs.add_image(img, "W1")
except Exception as e:
    print("logo:", e)

# ======================= PROYECTO =======================
pj = wb.create_sheet("PROYECTO", 0)
f(pj, "B1", "ANACAR — ANÁLISIS DE CARGAS DE TABLEROS ELÉCTRICOS", bold=True, size=18, color="006600")
for fila, et, nm, val in [(3, "Nombre del proyecto", "P_NOMBRE", ""), (4, "Proyecto N.º", "P_NUMERO", ""), (5, "Ubicación", "P_UBIC", ""), (6, "Fecha", "P_FECHA", ""), (7, "Elaboró", "P_ELABORO", ""), (8, "Revisión", "P_REV", "")]:
    f(pj, "B%d" % fila, et, size=10); f(pj, "C%d" % fila, val, fill=VERDE, borde=True); nombre(nm, "PROYECTO!$C$%d" % fila)
crit = [(3, "ΔV máx. alimentador (%)", "CVMAX_ALIM", 3, None), (4, "ΔV máx. total (%)", "CVMAX_TOT", 5, None), (5, "Desbalance máx. (%)", "DESB_MAX", 10, None), (6, "Long. máx. para Icc (m) (0 = sin límite)", "ICC_LMAX", 20, None),
        (7, "Marca por defecto", "MARCA_DEF", "Eaton", "L_MARCAS"), (8, "Reserva de espacios (%)", "RESERVA_ESP", 20, None), (9, "Ocupación (NEC)", "OCUPACION", "comercial", "L_OCUP"),
        (10, "Icc de la red (kA) (vacío = infinita)", "ICC_RED", None, None), (11, "Carga de tableros derivados", "CARGA_DERIV", "conectada", "L_CDERIV"), (12, "Aumentar calibre si no cumple 310.15", "AUTO_AMP", "SI", "L_SINO")]
for fila, et, nm, val, lista in crit:
    f(pj, "E%d" % fila, et, size=10); c = f(pj, "F%d" % fila, val, fill=VERDE, borde=True); nombre(nm, "PROYECTO!$F$%d" % fila)
    if lista:
        dv = DataValidation(type="list", formula1="=" + lista, allow_blank=True); pj.add_data_validation(dv); dv.add("F%d" % fila)
pj.column_dimensions["B"].width = 26; pj.column_dimensions["C"].width = 30; pj.column_dimensions["E"].width = 38; pj.column_dimensions["F"].width = 16
f(pj, "B15", "TABLEROS DEL PROYECTO (se actualiza con el botón; para editar un tablero abra su hoja MC)", bold=True, size=11)
cabp = ["Tablero", "Tipo", "Sistema", "Alimentado desde", "Long. (m)", "kVA totales", "kVA demandados", "Protección (A)", "V bornes", "ΔV acum. (%)", "Icc (kA)", "Tablero (catálogo)", "Avisos", "Orden resumen", "Excluir (X)"]
for j, t in enumerate(cabp):
    c = pj.cell(16, 2 + j, t); c.font = Font(name=FN, bold=True, size=9); c.fill = GRIS; c.alignment = CENT; c.border = BORDE
for col, w in zip("BCDEFGHIJKLMNOP", [18, 6, 10, 18, 9, 11, 11, 11, 10, 10, 9, 22, 60, 9, 9]): pj.column_dimensions[col].width = w
nombre("LISTA_TABLEROS", "PROYECTO!$B$17:$B$116")
pj.freeze_panes = "B17"

# ======================= PREVISTA =======================
pv = wb.create_sheet("PREVISTA", 1)
f(pv, "B1", "PREVISTA (diseño esquemático) — carga por tipo de uso y m²", bold=True, size=16, color="006600")
f(pv, "B2", "Iluminación: NEC 2020 tabla 220.12; tomas y climatización: densidades estimadas (editar en CAT_USOS como administrador). Vacío = valor del tipo de uso.", size=9)
cabv = ["Área", "Tipo de uso", "m²", "Ilum. VA/m²", "Tomas VA/m²", "Clima VA/m²", "Otros kVA", "Tablero", "Long. (m)", "Ilum. efectiva", "Tomas efectiva", "Clima efectiva", "kVA ilum.", "kVA tomas", "kVA clima+otros", "kVA total", "Circuitos generados"]
for j, t in enumerate(cabv):
    c = pv.cell(4, 2 + j, t); c.font = Font(name=FN, bold=True, size=9); c.fill = GRIS; c.alignment = CENT; c.border = BORDE
for r in range(5, 55):
    for col in "BCDEFGHIJ": pv["%s%d" % (col, r)].fill = VERDE; pv["%s%d" % (col, r)].border = BORDE
    pv["K%d" % r] = '=IF($C{r}="","",IF(ISNUMBER($E{r}),$E{r},IFERROR(VLOOKUP($C{r},T_USOS,2,FALSE),0)))'.format(r=r)
    pv["L%d" % r] = '=IF($C{r}="","",IF(ISNUMBER($F{r}),$F{r},IFERROR(VLOOKUP($C{r},T_USOS,3,FALSE),0)))'.format(r=r)
    pv["M%d" % r] = '=IF($C{r}="","",IF(ISNUMBER($G{r}),$G{r},IFERROR(VLOOKUP($C{r},T_USOS,4,FALSE),0)))'.format(r=r)
    pv["N%d" % r] = '=IF($C{r}="",0,N($D{r})*$K{r}/1000)'.format(r=r); pv["O%d" % r] = '=IF($C{r}="",0,N($D{r})*$L{r}/1000)'.format(r=r)
    pv["P%d" % r] = '=IF($C{r}="",0,N($D{r})*$M{r}/1000+N($H{r}))'.format(r=r); pv["Q%d" % r] = '=N{r}+O{r}+P{r}'.format(r=r)
    for col in "NOPQ": pv["%s%d" % (col, r)].number_format = "0.00"
dv = DataValidation(type="list", formula1="=L_USOS", allow_blank=True); pv.add_data_validation(dv); dv.add("C5:C54")
dv = DataValidation(type="list", formula1="=LISTA_TABLEROS", allow_blank=True); pv.add_data_validation(dv); dv.add("I5:I54")
pv["P56"] = "TOTAL kVA"; pv["Q56"] = "=SUM(Q5:Q54)"; pv["P56"].font = Font(name=FN, bold=True); pv["Q56"].font = Font(name=FN, bold=True); pv["Q56"].number_format = "0.00"
for col, w in zip("BCDEFGHIJKLMNOPQR", [24, 26, 9, 10, 10, 10, 9, 16, 9, 10, 10, 10, 10, 10, 12, 10, 12]): pv.column_dimensions[col].width = w

# ======================= REVIT, CAMBIOS, TABLAS RESUMEN =======================
rv = wb.create_sheet("REVIT", 2)
f(rv, "B1", "IMPORTAR / ACTUALIZAR DESDE LA TABLA DE CIRCUITOS DE REVIT", bold=True, size=16, color="006600")
f(rv, "B2", "Opción 1: botón 'Importar archivo' (.txt/.csv exportado de Revit o el Excel 'Circuitos revit-excel'). Opción 2: pegue la tabla desde B5 (con encabezados) y use 'Importar lo pegado'.", size=9)
f(rv, "B3", "Columnas esperadas: Voltage, True Load, Panel, Circuit Number, Number of Poles, Load Name, Length. Los tableros que ya existen se ACTUALIZAN (carga, longitud y nombre) conservando los demás ajustes.", size=9)
cm = wb.create_sheet("CAMBIOS REVIT", 3)
for j, t in enumerate(["Fecha", "Panel (Revit)", "Carga", "Circuito en Revit", "Circuito nuevo", "Origen", "Aplicado (X)"]):
    c = cm.cell(1, 1 + j, t); c.font = Font(name=FN, bold=True); c.fill = GRIS; c.border = BORDE
for col, w in zip("ABCDEFG", [12, 14, 40, 16, 16, 14, 12]): cm.column_dimensions[col].width = w

tr = wb.create_sheet("TABLA RESUMEN", 4)
ar = {"A": 4, "B": 35.8, "C": 30.2, "D": 18.8, "E": 18.9, "F": 15.5, "G": 15.8, "H": 15.1, "I": 15, "J": 9, "K": 15.5, "L": 13.9, "M": 13.9, "N": 9, "O": 13.8, "P": 18.2, "Q": 16.8, "R": 10, "S": 17.8, "T": 12, "U": 14, "V": 14, "W": 14,
      "X": 19.1, "Y": 18.2, "Z": 32.1, "AA": 23.5, "AB": 16.5, "AC": 15, "AD": 10, "AE": 10, "AF": 16, "AG": 20.2, "AH": 21.8, "AI": 34.9, "AJ": 21.8, "AK": 12, "AL": 14, "AM": 12, "AN": 21.5, "AO": 30.9, "AP": 43, "AQ": 16.2, "AR": 21.8, "AS": 14, "AT": 10, "AU": 12}
for k, w in ar.items(): tr.column_dimensions[k].width = w
def T_(y, a, b, val, est, y2=None):
    y2 = y2 or y
    for yy in range(y, y2 + 1):
        for x in range(CI(a), CI(b) + 1):
            c = tr.cell(yy, x); c.border = BORDEM if est == "tit" else BORDE
            if est == "thg": c.fill = GRIS
    c = tr["%s%d" % (a, y)]; c.value = val
    c.font = Font(name=FN, size=40 if est == "tit" else 12, bold=True, color="006600" if est == "tit" else None)
    c.alignment = CENT
    if a != b or y2 != y: tr.merge_cells("%s%d:%s%d" % (a, y, b, y2))
T_(1, "B", "N", "TABLA RESUMEN - TABLEROS ELÉCTRICOS", "tit", 8); T_(1, "O", "Y", None, "tit", 8)
T_(1, "Z", "AG", "DATOS DEL TABLERO", "tit", 8); T_(1, "AH", "AN", "DATOS DEL SUPRESOR", "tit", 8); T_(1, "AO", "AU", "DATOS INTERRUPTOR PRINCIPAL", "tit", 8)
for L_, t in [("B", "Tablero / Equipo"), ("C", "Alimentado desde"), ("D", "kVA Totales"), ("E", "kVA Demandados"), ("F", "Factor demanda"), ("G", "Factor diversidad"), ("H", "Factor potencia"),
              ("T", "Longitud (m)"), ("U", "Voltaje bornes (V)"), ("V", "Caída de voltaje total (V)"), ("W", "Caída de voltaje total (%)"), ("X", "Corriente cortocircuito disponible (kA)"), ("Y", "Capacidad de breaker (A)"),
              ("Z", "Tablero / Equipo"), ("AA", "Modelo de referencia"), ("AB", "Fabricante"), ("AF", "Espacios"), ("AG", "Montaje"), ("AH", "Fabricante"), ("AI", "Modelo de referencia"), ("AJ", "Montaje"),
              ("AM", "Voltaje (V)"), ("AN", "Fases"), ("AO", "Fabricante"), ("AP", "Modelo de referencia"), ("AQ", "Marco (A)"), ("AR", "Amperios (A)"), ("AS", "Tipo de unidad"), ("AT", "# polos"), ("AU", "SCCR (kA)")]:
    T_(9, L_, L_, t, "thg", 10)
T_(9, "I", "P", "Calibres conductores alimentador", "thg")
for a, b, t in [("I", "J", "Fases (AWG)"), ("K", "L", "Neutro (AWG)"), ("M", "N", "Tierra (AWG)"), ("O", "O", "Material"), ("P", "P", "Aislamiento")]: T_(10, a, b, t, "thg")
T_(9, "Q", "S", "Tubería", "thg"); [T_(10, L_, L_, t, "thg") for L_, t in [("Q", "Cantidad"), ("R", "(Ø mm)"), ("S", "Tipo")]]
T_(9, "AC", "AE", "Barras", "thg"); [T_(10, L_, L_, t, "thg") for L_, t in [("AC", "Fase"), ("AD", "Neutro"), ("AE", "Tierra")]]
T_(9, "AK", "AL", "Capacidad de supresión (kA)", "thg"); T_(10, "AK", "AK", "L-L", "thg"); T_(10, "AL", "AL", "L-N / L-T / N-T", "thg")
tr.row_dimensions[9].height = 23; tr.row_dimensions[10].height = 44
tr.sheet_view.zoomScale = 60; tr.sheet_view.showGridLines = False
try:
    img = XLImage(LOGO); img.width, img.height = 420, 150; tr.add_image(img, "R2")
except Exception: pass
for nm, titulo in [("RESUMEN VERTICAL", "TABLA RESUMEN - TABLEROS ELÉCTRICOS (VERTICAL)"), ("RESUMEN DU", "TABLA RESUMEN DU - TABLEROS ELÉCTRICOS")]:
    w2 = wb.create_sheet(nm, 5)
    f(w2, "B1", titulo, bold=True, size=24, color="006600"); w2.column_dimensions["B"].width = 42
    w2.sheet_view.showGridLines = False

# ======================= INICIO =======================
ini = wb.create_sheet("INICIO", 0)
f(ini, "B2", "ANACAR — Análisis de cargas de tableros eléctricos", bold=True, size=22, color="006600")
f(ini, "B3", "Versión Excel con macros de la herramienta web (https://oquirosme-sys.github.io/Herramienta-ANACAR/). Base: ANACAR 2026.1 · NEC 2020.", size=10)
pasos = ["1. Llene los datos y criterios en la hoja PROYECTO.",
         "2. Cree los tableros con 'Nuevo tablero' o impórtelos/actualícelos desde la tabla de circuitos de Revit (hoja REVIT).",
         "3. En cada hoja 'MC <tablero>' (memoria de cálculo) llene los datos en las celdas verdes. 'Alimentado desde' conecta el tablero a su padre y crea su circuito.",
         "4. Las hojas '3F/1F <tablero>' muestran el cuadro de cargas con el diseño del Excel original; 'TABLA RESUMEN', 'RESUMEN VERTICAL' y 'RESUMEN DU' son para los planos de Revit.",
         "5. 'Actualizar tablas' rehace la lista de tableros, las tablas resumen y los nombres TABLERO_<nombre> para vincular en Revit.",
         "6. 'Autobalancear' reubica circuitos del tablero activo y anota los cambios en la hoja CAMBIOS REVIT.",
         "7. 'Exportar DXF' genera el diagrama unifilar para AutoCAD; 'Exportar PDF' imprime los cuadros, tablas y memorias.",
         "8. Catálogos (marcas, tableros, breakers, supresores, cargas, conductores, reglas): botón 'Modo administrador' (contraseña inicial: sinergia-admin).",
         "Habilite las macros al abrir el archivo. Celdas verdes = datos; celdas vacías = valor automático."]
for i, t in enumerate(pasos): f(ini, "B%d" % (6 + i), t, size=11)
ini.column_dimensions["B"].width = 140
ini.sheet_view.showGridLines = False

# orden y visibilidad
for nm in wb.sheetnames:
    if nm.startswith("CAT_"): wb[nm].sheet_state = "hidden"
wb["MC_MACHOTE"].sheet_state = "hidden"; wb["VISTA_MACHOTE"].sheet_state = "hidden"
del wb["Sheet"]
orden = ["INICIO", "PROYECTO", "PREVISTA", "REVIT", "CAMBIOS REVIT", "TABLA RESUMEN", "RESUMEN VERTICAL", "RESUMEN DU", "MC_MACHOTE", "VISTA_MACHOTE"]
wb._sheets = [wb[n] for n in orden] + [s for s in wb._sheets if s.title not in orden]
wb.save(OUT)
print("ok", OUT, len(wb.sheetnames))
